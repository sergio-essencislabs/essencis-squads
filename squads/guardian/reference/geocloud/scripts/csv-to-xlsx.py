#!/usr/bin/env python3
"""
Converte um CSV "multi-aba" (várias abas concatenadas em um único arquivo,
separadas por uma linha marcadora ``===== SHEET: <Nome> =====``) em um
arquivo .xlsx com uma planilha por aba.

Este é o mesmo padrão já usado historicamente em
`GeoCloud/Miscelaneous/Account_resumo_estrutural.csv` (abas Atributos,
Metodos_Back, Permissões dentro de um único .csv). O CSV continua sendo a
fonte canônica (diffável em git); o .xlsx é apenas um espelho gerado para
visualização direta em Excel/LibreOffice.

Uso:
    python csv-to-xlsx.py <entrada.csv> <saida.xlsx>

Requer: openpyxl (`pip install openpyxl`).
"""
import csv
import sys
from pathlib import Path

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter
except ImportError:
    sys.exit("Erro: pacote 'openpyxl' não encontrado. Instale com: pip install openpyxl")

SHEET_MARKER_PREFIX = "===== SHEET:"
SHEET_MARKER_SUFFIX = "====="
MAX_COL_WIDTH = 60
MIN_COL_WIDTH = 8


def parse_sheet_name(marker_line: str) -> str:
    name = marker_line.strip()
    name = name[len(SHEET_MARKER_PREFIX):].strip()
    if name.endswith(SHEET_MARKER_SUFFIX):
        name = name[: -len(SHEET_MARKER_SUFFIX)].strip()
    # Excel sheet name limits: 31 chars, no []:*?/\\
    for ch in "[]:*?/\\":
        name = name.replace(ch, "-")
    return name[:31] if name else "Sheet"


def split_into_sheets(csv_path: Path):
    sheets = []  # list of (name, list[list[str]])
    current_name = None
    current_rows = []

    with csv_path.open("r", encoding="utf-8-sig", newline="") as f:
        reader = csv.reader(f)
        for row in reader:
            is_marker = len(row) >= 1 and row[0].strip().startswith(SHEET_MARKER_PREFIX)
            if is_marker:
                if current_name is not None:
                    sheets.append((current_name, current_rows))
                current_name = parse_sheet_name(row[0])
                current_rows = []
            else:
                if current_name is None:
                    # Conteúdo antes do primeiro marcador — agrupa em uma aba default.
                    current_name = "Sheet1"
                current_rows.append(row)
        if current_name is not None:
            sheets.append((current_name, current_rows))

    return sheets


def autosize_columns(ws):
    widths = {}
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is None:
                continue
            length = len(str(cell.value))
            col = cell.column_letter
            widths[col] = max(widths.get(col, MIN_COL_WIDTH), min(length + 2, MAX_COL_WIDTH))
    for col, width in widths.items():
        ws.column_dimensions[col].width = width


def build_workbook(sheets, header_bold=True):
    wb = Workbook()
    wb.remove(wb.active)

    used_names = set()
    for name, rows in sheets:
        safe_name = name
        suffix = 2
        while safe_name in used_names:
            safe_name = f"{name[:28]}_{suffix}"
            suffix += 1
        used_names.add(safe_name)

        ws = wb.create_sheet(title=safe_name)
        for r_idx, row in enumerate(rows, start=1):
            for c_idx, value in enumerate(row, start=1):
                ws.cell(row=r_idx, column=c_idx, value=value if value != "" else None)
        if header_bold and rows:
            for c_idx in range(1, len(rows[0]) + 1):
                ws.cell(row=1, column=c_idx).font = Font(bold=True)
        ws.freeze_panes = "A2"
        autosize_columns(ws)

    return wb


def main():
    if len(sys.argv) != 3:
        sys.exit("Uso: python csv-to-xlsx.py <entrada.csv> <saida.xlsx>")

    csv_path = Path(sys.argv[1])
    xlsx_path = Path(sys.argv[2])

    if not csv_path.exists():
        sys.exit(f"Erro: arquivo não encontrado: {csv_path}")

    sheets = split_into_sheets(csv_path)
    if not sheets:
        sys.exit(f"Erro: nenhuma aba encontrada em {csv_path} (esperado marcador '{SHEET_MARKER_PREFIX} <Nome> {SHEET_MARKER_SUFFIX}').")

    wb = build_workbook(sheets)
    xlsx_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(xlsx_path)

    print(f"OK: {xlsx_path}")
    for name, rows in sheets:
        print(f"  - aba '{name}': {len(rows)} linhas")


if __name__ == "__main__":
    main()
